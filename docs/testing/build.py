import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
sys.path.insert(0, ".")
from cases import C  # noqa: E402

out = sys.argv[1]
wb = Workbook()
thin = Side(style="thin", color="D0D4DA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="1F3A5F")
hdr_font = Font(bold=True, color="FFFFFF")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="top", wrap_text=True)

# --- 说明 ---
ws0 = wb.active
ws0.title = "说明"
intro = [
    ("Workflow System 功能测试表", True),
    ("版本：v0.1.0（main）　编写日期：2026-10-12", False),
    ("", False),
    ("怎么用", True),
    ("1. 在\"测试用例\"页从上往下测，P0 先测（必须通过才能上线），P1 次之，P2 有空再测。", False),
    ("2. 每条按\"操作步骤\"做，对照\"预期结果\"，在\"测试结果\"列下拉选：通过 / 不通过 / 阻塞 / 跳过。", False),
    ("3. 不通过时在\"问题/备注\"写清楚实际看到了什么（最好附截图文件名）。", False),
    ("4. \"汇总\"页会自动统计进度和通过率。", False),
    ("", False),
    ("\"老系统要求\"和\"对照\"两列", True),
    ("老系统要求：摘自 new-system-doc.md 第一部分，§ 是章节号。这是最低标准，新系统必须做到。", False),
    ("同等：与老系统一致。　增强：做到老系统的要求，并且更多。　新增：老系统没有。", False),
    ("差异：之前已决定有意改变做法（例如客户字段只存一份、删除改为回收站），测试时确认新做法业务上能接受。", False),
    ("待确认：代码里没看到老系统的这条规则，或行为可能不同。这几条请重点测，结果告诉我，再决定是否补。", False),
    ("不适用：老系统也只是占位。", False),
    ("", False),
    ("测试前准备", True),
    ("1. 用管理员账号登录。在 Users 里建 4 个测试用户：", False),
    ("　　A：两条线都是 member（日常操作）　B：只有 Accounting member　C：Accounting viewer　D：任意角色（用于停用测试）", False),
    ("2. 只用虚构的测试数据（如 \"Test Client A Pty Ltd\"），不要用真实客户。", False),
    ("3. 冲突、实时刷新、@ 通知需要两个浏览器同时登录（如 Chrome 普通窗口 + 无痕窗口）。", False),
    ("4. 日期类自动化每小时自动跑一次；想立即看到结果，在服务器执行：", False),
    ("　　docker compose -f deploy/docker-compose.yml exec worker python -m app.cli run-job automations.date_rules", False),
    ("5. 数据迁移用 D:\\Projects\\Workflow System\\csv sample 里的两个导出文件。", False),
]
for i, (t, bold) in enumerate(intro, 1):
    c = ws0.cell(row=i, column=1, value=t)
    c.font = Font(bold=bold, size=14 if i == 1 else 11)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws0.column_dimensions["A"].width = 120

# --- 测试用例 ---
ws = wb.create_sheet("测试用例")
headers = ["编号", "模块", "功能点", "优先级", "前置条件", "操作步骤", "预期结果", "老系统要求（最低标准）",
           "对照", "测试结果", "测试日期", "测试人", "问题/备注"]
widths = [8, 16, 20, 7, 28, 44, 44, 34, 9, 10, 11, 9, 32]
for j, (h, w) in enumerate(zip(headers, widths), 1):
    c = ws.cell(row=1, column=j, value=h)
    c.fill, c.font, c.alignment, c.border = hdr_fill, hdr_font, center, border
    ws.column_dimensions[get_column_letter(j)].width = w

prefix, counters = {}, {}
mods = []
for m, *_ in C:
    if m not in mods:
        mods.append(m)
for m in mods:
    prefix[m] = f"{mods.index(m) + 1:02d}"
cmp_fill = {"差异": "FFF4E5", "待确认": "FDE2E1", "新增": "E8F4FD", "增强": "EAF7EE"}
for i, (m, feat, pri, pre, steps, exp, old, cmp) in enumerate(C, 2):
    counters[m] = counters.get(m, 0) + 1
    row = [f"{prefix[m]}-{counters[m]:02d}", m, feat, pri, pre, steps, exp, old, cmp, None, None, None, None]
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.alignment = center if j in (1, 4, 9, 10, 11, 12) else wrap
        c.border = border
    if cmp in cmp_fill:
        ws.cell(row=i, column=9).fill = PatternFill("solid", fgColor=cmp_fill[cmp])
    if pri == "P0":
        ws.cell(row=i, column=4).font = Font(bold=True, color="C0392B")
last = len(C) + 1
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:M{last}"
dv = DataValidation(type="list", formula1='"通过,不通过,阻塞,跳过"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"J2:J{last}")
dvd = DataValidation(type="date", allow_blank=True)
ws.add_data_validation(dvd)
dvd.add(f"K2:K{last}")
for col in ("K",):
    for r in range(2, last + 1):
        ws[f"{col}{r}"].number_format = "yyyy-mm-dd"
rng = f"J2:J{last}"
for val, color in [("通过", "C6EFCE"), ("不通过", "FFC7CE"), ("阻塞", "FFEB9C"), ("跳过", "E7E6E6")]:
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=color)))

# --- 汇总 ---
s = wb.create_sheet("汇总", 1)
sh = ["模块", "用例数", "P0", "已测", "通过", "不通过", "阻塞", "跳过", "未测", "通过率（已测）"]
for j, h in enumerate(sh, 1):
    c = s.cell(row=1, column=j, value=h)
    c.fill, c.font, c.alignment, c.border = hdr_fill, hdr_font, center, border
    s.column_dimensions[get_column_letter(j)].width = 14 if j > 1 else 24
R = f"测试用例!$J$2:$J${last}"
B = f"测试用例!$B$2:$B${last}"
P = f"测试用例!$D$2:$D${last}"
for i, m in enumerate(mods + ["合计"], 2):
    s.cell(row=i, column=1, value=m)
    if m == "合计":
        for j in range(2, 10):
            L = get_column_letter(j)
            s.cell(row=i, column=j, value=f"=SUM({L}2:{L}{i - 1})")
        for j in range(1, 11):
            s.cell(row=i, column=j).font = Font(bold=True)
    else:
        a = f"$A{i}"
        s.cell(row=i, column=2, value=f'=COUNTIF({B},{a})')
        s.cell(row=i, column=3, value=f'=COUNTIFS({B},{a},{P},"P0")')
        s.cell(row=i, column=4, value=f"=E{i}+F{i}+G{i}+H{i}")
        for j, v in zip(range(5, 9), ["通过", "不通过", "阻塞", "跳过"]):
            s.cell(row=i, column=j, value=f'=COUNTIFS({B},{a},{R},"{v}")')
        s.cell(row=i, column=9, value=f"=B{i}-D{i}")
    s.cell(row=i, column=10, value=f'=IF(D{i}-H{i}=0,"-",E{i}/(D{i}-H{i}))')
    s.cell(row=i, column=10).number_format = "0%"
    for j in range(1, 11):
        s.cell(row=i, column=j).border = border
        if j > 1:
            s.cell(row=i, column=j).alignment = Alignment(horizontal="center")
n = len(mods) + 3
s.cell(row=n + 1, column=1, value="P0 未通过（不通过 + 阻塞）").font = Font(bold=True)
s.cell(row=n + 1, column=2, value=f'=COUNTIFS({P},"P0",{R},"不通过")+COUNTIFS({P},"P0",{R},"阻塞")')
s.cell(row=n + 2, column=1, value="上线条件：P0 全部通过；\"待确认\"项都有结论").font = Font(italic=True, color="666666")
wb.save(out)
print(len(C), "cases,", len(mods), "modules")
